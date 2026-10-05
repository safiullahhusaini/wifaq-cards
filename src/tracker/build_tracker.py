import datetime as dt
import os
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule, FormulaRule, DataBarRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

TEST = "--test" in sys.argv
OUT = sys.argv[1] if len(sys.argv) > 1 and not sys.argv[1].startswith("--") else "Wifaq_1448_Tracker.xlsx"

FONT = "Arial"
GREEN = "1F5C4A"
LIGHT = "E8F1EC"
YELLOW = "FFF4C2"
GREY = "F2F2F2"

f_title = Font(name=FONT, size=16, bold=True, color=GREEN)
f_sub = Font(name=FONT, size=10, italic=True, color="555555")
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_label = Font(name=FONT, size=10, bold=True, color="333333")
f_body = Font(name=FONT, size=10, color="000000")
f_input = Font(name=FONT, size=10, color="0000FF")
f_note = Font(name=FONT, size=9, italic=True, color="666666")

fill_head = PatternFill("solid", fgColor=GREEN)
fill_input = PatternFill("solid", fgColor=YELLOW)
fill_light = PatternFill("solid", fgColor=LIGHT)
fill_grey = PatternFill("solid", fgColor=GREY)

thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center", wrap_text=True)
DATE_FMT = "d mmm yyyy"

wb = Workbook()

# ---------------------------------------------------------------- lists sheet
ls = wb.active
ls.title = "فہرستیں"
books = [
    # paper, name, marks, min/page, weekly hours
    ("پرچہ 1", "التبیان فی علوم القرآن", 30, 6, 1.0),
    ("پرچہ 1", "تیسیر مصطلح الحدیث", 25, 6, 0.5),
    ("پرچہ 1", "شرح نخبۃ الفکر", 25, 8, 0.5),
    ("پرچہ 1", "آئینہ قادیانیت", 20, 3, 0.5),
    ("پرچہ 2", "بیضاوی (ربع پارہ اول)", 100, 20, 2.5),
    ("پرچہ 3", "مشکوٰۃ المصابیح جلد اول", 100, 8, 3.75),
    ("پرچہ 4", "مشکوٰۃ المصابیح جلد دوم", 100, 8, 3.75),
    ("پرچہ 5", "ہدایہ جلد ثالث", 100, 12, 4.5),
    ("پرچہ 6", "ہدایہ جلد رابع", 100, 12, 4.5),
]
methods = ["تمہیدی کارڈ + عبارت", "ریکارڈنگ", "کلاس", "شرح", "سابقہ پرچہ", "دہرائی"]
ls["A1"] = "کتابیں"
ls["B1"] = "طریقہ"
for c in ("A1", "B1"):
    ls[c].font = f_label
for i, b in enumerate(books):
    ls.cell(row=2 + i, column=1, value=b[1]).font = f_body
for i, m in enumerate(methods):
    ls.cell(row=2 + i, column=2, value=m).font = f_body
ls.column_dimensions["A"].width = 30
ls.column_dimensions["B"].width = 22
ls.sheet_view.rightToLeft = True
BOOK_LIST = "'فہرستیں'!$A$2:$A$10"
METHOD_LIST = "'فہرستیں'!$B$2:$B$7"


def head_row(ws, row, headers, height=32):
    for col, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=col, value=h)
        c.font = f_head
        c.fill = fill_head
        c.alignment = center
        c.border = box
    ws.row_dimensions[row].height = height


def inp(cell, value=None, fmt=None):
    if value is not None:
        cell.value = value
    cell.font = f_input
    cell.fill = fill_input
    cell.border = box
    cell.alignment = center
    if fmt:
        cell.number_format = fmt


def calc(cell, value, fmt=None, bold=False):
    cell.value = value
    cell.font = Font(name=FONT, size=10, bold=bold)
    cell.border = box
    cell.alignment = center
    if fmt:
        cell.number_format = fmt


# ---------------------------------------------------------------- main sheet
q = wb.create_sheet("کوٹہ", 0)
q.sheet_view.rightToLeft = True
q["A1"] = "عالمیہ سال اول (موقوف علیہ) — وفاق 1448ھ کی تیاری"
q["A1"].font = f_title
q.merge_cells("A1:W1")
q["A2"] = ("امتحان: یکم تا 6 شعبان 1448ھ = ہفتہ 9 تا جمعرات 14 جنوری 2027 (وفاق المدارس کا اعلان)۔ "
           "پیلے خانے آپ بھریں؛ باقی سب خودکار ہے۔ روزانہ کا کام 'لاگ' میں لکھیں۔")
q["A2"].font = f_sub
q.merge_cells("A2:W2")

labels = [
    ("B4", "آج", "C4", "=TODAY()", DATE_FMT, False),
    ("B5", "امتحان شروع", "C5", dt.date(2027, 1, 9), DATE_FMT, True),
    ("B6", "امتحان تک دن", "C6", "=C5-C4", "0", False),
    ("B7", "پہلا دور ختم", "C7", dt.date(2026, 12, 13), DATE_FMT, True),
    ("B8", "پہلے دور کے باقی دن", "C8", "=IF(C7>=C4,C7-C4+1,0)", "0", False),
    ("E4", "روزانہ گھنٹے", "F4", 3.5, "0.0", True),
    ("E5", "ہفتہ وار دستیاب گھنٹے", "F5", "=F4*7", "0.0", False),
    ("E6", "ہفتہ وار مختص گھنٹے", "F6", "=SUM(Q11:Q21)", "0.0", False),
    ("E7", "درجہ B وقت کا تناسب", "F7", 0.35, "0%", True),
    ("E8", "درجہ C وقت کا تناسب", "F8", 0.10, "0%", True),
]
for lc, lt, vc, vv, fmt, is_input in labels:
    q[lc] = lt
    q[lc].font = f_label
    q[lc].alignment = right
    if is_input:
        inp(q[vc], vv, fmt)
    else:
        calc(q[vc], vv, fmt, bold=True)
q["C5"].comment = Comment("ماخذ: وفاق المدارس العربیہ کا اعلان، یکم تا 6 شعبان 1448ھ مطابق 9 تا 14 جنوری 2027", "Claude")
q["C7"].comment = Comment("پہلا دور 13 دسمبر تک؛ اس کے بعد 14 تا 31 دسمبر دہرائی اور 1 تا 8 جنوری آخری دہرائی", "Claude")
q["F7"].comment = Comment("درجہ B (صرف تمہیدی کارڈ + عبارت پر نظر) میں مکمل مطالعے کا تقریباً کتنا وقت لگتا ہے۔ پہلے ہفتے کے بعد اپنے تجربے سے بدلیں۔", "Claude")
q["F8"].comment = Comment("درجہ C (صرف عنوانات اور مسائل کی فہرست) کا وقت، مکمل مطالعے کے مقابلے میں۔", "Claude")

WS = "'ہفتہ وار نقشہ'"
side = [
    ("H4", "موجودہ مرحلہ", "I4",
     f'=IFERROR(INDEX({WS}!$D$5:$D$32,MATCH($C$4,{WS}!$B$5:$B$32,1)),"ابھی شروع نہیں ہوا")'),
    ("H5", "اس ہفتے / دن کا کام", "I5",
     f'=IFERROR(INDEX({WS}!$E$5:$E$32,MATCH($C$4,{WS}!$B$5:$B$32,1)),"")'),
    ("H6", "آج: بلاک 1 (90 منٹ)", "I6",
     f'=IFERROR(VLOOKUP(WEEKDAY($C$4),{WS}!$A$37:$F$43,3,FALSE),"")'),
    ("H7", "آج: بلاک 2 (75 منٹ)", "I7",
     f'=IFERROR(VLOOKUP(WEEKDAY($C$4),{WS}!$A$37:$F$43,4,FALSE),"")'),
    ("H8", "آج: بلاک 3 (30 منٹ)", "I8",
     f'=IFERROR(VLOOKUP(WEEKDAY($C$4),{WS}!$A$37:$F$43,5,FALSE),"")'),
]
for lc, lt, vc, vv in side:
    q[lc] = lt
    q[lc].font = f_label
    q[lc].alignment = right
    calc(q[vc], vv)
    q[vc].alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    q.merge_cells(f"{vc}:M{vc[1:]}")
for r in range(4, 9):
    q.row_dimensions[r].height = 36
q["O4"] = "کل پیش رفت (پہلا دور)"
q["O4"].font = f_label
q["O4"].alignment = right
calc(q["P4"], '=IFERROR(SUM(K11:K19)/SUM(G11:G19),0)', "0%", bold=True)
q["O5"] = "آج تک کا وقت (گھنٹے)"
q["O5"].font = f_label
q["O5"].alignment = right
calc(q["P5"], "=SUM('لاگ'!$F$6:$F$1000)/60", "0.0", bold=True)
q["O6"] = "حفظ: آج دہرائی"
q["O6"].font = f_label
q["O6"].alignment = right
calc(q["P6"], "=COUNTIF('حفظ احادیث'!$K$6:$K$43,\"آج دہرائیں\")", "0", bold=True)

headers = ["#", "پرچہ", "کتاب", "نمبر", "پہلا صفحہ", "آخری صفحہ", "کل صفحات",
           "درجہ A صفحات", "درجہ B صفحات", "درجہ C صفحات", "مکمل (پہلا دور)", "باقی",
           "فیصد", "منٹ فی صفحہ (A)", "اصل منٹ فی صفحہ", "درکار گھنٹے", "ہفتہ وار گھنٹے",
           "دستیاب گھنٹے", "حالت", "ہفتہ وار ہدف (صفحات)", "دہرائی (صفحات)", "استاد / ریکارڈنگ",
           "اردو شرح"]
head_row(q, 10, headers, height=44)

LOG_B = "'لاگ'!$B$6:$B$1000"
LOG_E = "'لاگ'!$E$6:$E$1000"
LOG_F = "'لاگ'!$F$6:$F$1000"
LOG_G = "'لاگ'!$G$6:$G$1000"

for i, (paper, name, marks, mpp, wh) in enumerate(books):
    r = 11 + i
    calc(q.cell(r, 1), i + 1)
    calc(q.cell(r, 2), paper)
    c = q.cell(r, 3, name)
    c.font = Font(name=FONT, size=10, bold=True)
    c.border = box
    c.alignment = right
    calc(q.cell(r, 4), marks)
    inp(q.cell(r, 5))
    inp(q.cell(r, 6))
    calc(q.cell(r, 7), f'=IF(AND(ISNUMBER(E{r}),ISNUMBER(F{r})),IF(F{r}>=E{r},F{r}-E{r}+1,""),"")')
    inp(q.cell(r, 8))
    inp(q.cell(r, 9))
    calc(q.cell(r, 10), f'=IF(G{r}="","",MAX(0,G{r}-IF(H{r}="",G{r},H{r})-IF(I{r}="",0,I{r})))')
    calc(q.cell(r, 11), f'=SUMIFS({LOG_E},{LOG_B},$C{r},{LOG_G},"<>دہرائی")')
    calc(q.cell(r, 12), f'=IF(G{r}="","",MAX(0,G{r}-K{r}))')
    calc(q.cell(r, 13), f'=IF(G{r}="","",MIN(1,K{r}/G{r}))', "0%")
    inp(q.cell(r, 14), mpp, "0")
    calc(q.cell(r, 15), f'=IF(K{r}=0,"",SUMIFS({LOG_F},{LOG_B},$C{r},{LOG_G},"<>دہرائی")/K{r})', "0.0")
    calc(q.cell(r, 16),
         f'=IF(G{r}="","",(IF(H{r}="",G{r},H{r})*N{r}+IF(I{r}="",0,I{r})*N{r}*$F$7+J{r}*N{r}*$F$8)/60*L{r}/G{r})',
         "0")
    inp(q.cell(r, 17), wh, "0.00")
    calc(q.cell(r, 18), f'=Q{r}*$C$8/7', "0")
    calc(q.cell(r, 19),
         f'=IF(G{r}="","صفحات درج کریں",IF(L{r}=0,"مکمل ✓",IF(R{r}>=P{r},"ٹھیک","پیچھے: "&TEXT(P{r}-R{r},"0")&" گھنٹے کم")))')
    calc(q.cell(r, 20), f'=IF(OR(G{r}="",$C$8=0),"",L{r}/($C$8/7))', "0")
    calc(q.cell(r, 21), f'=SUMIFS({LOG_E},{LOG_B},$C{r},{LOG_G},"دہرائی")')
    inp(q.cell(r, 22))
    inp(q.cell(r, 23))

# hifz row
r = 20
calc(q.cell(r, 1), 10)
calc(q.cell(r, 2), "پرچہ 3")
c = q.cell(r, 3, "حفظ احادیث (38)")
c.font = Font(name=FONT, size=10, bold=True)
c.border = box
c.alignment = right
calc(q.cell(r, 4), 15)
for col in (5, 6, 8, 9, 10, 14, 15, 16, 18, 21, 22, 23):
    calc(q.cell(r, col), "")
    q.cell(r, col).fill = fill_grey
calc(q.cell(r, 7), 38)
calc(q.cell(r, 11), "=COUNT('حفظ احادیث'!$F$6:$F$43)")
calc(q.cell(r, 12), "=MAX(0,G20-K20)")
calc(q.cell(r, 13), "=MIN(1,K20/G20)", "0%")
inp(q.cell(r, 17), 1.75, "0.00")
calc(q.cell(r, 19), '=IF(L20=0,"مکمل ✓",L20&" احادیث باقی")')
calc(q.cell(r, 20), '=IF($C$8=0,"",L20/($C$8/7))', "0.0")

# weekly review row
r = 21
calc(q.cell(r, 1), 11)
calc(q.cell(r, 2), "سب")
c = q.cell(r, 3, "ہفتہ وار جائزہ / سابقہ پرچے")
c.font = Font(name=FONT, size=10, bold=True)
c.border = box
c.alignment = right
for col in list(range(4, 17)) + list(range(18, 24)):
    calc(q.cell(r, col), "")
    q.cell(r, col).fill = fill_grey
inp(q.cell(r, 17), 1.25, "0.00")

q.row_dimensions[22].height = 8
q["A23"] = "کیسے استعمال کریں"
q["A23"].font = f_label
notes = [
    "1) اپنی بشریٰ کی کتابوں سے ہر کتاب کا پہلا اور آخری صفحہ (نصاب کے مطابق) پیلے خانوں میں لکھیں۔",
    "2) درجہ A = پورا طریقہ (تمہید + عبارت + مشکل الفاظ)؛ درجہ B = صرف تمہیدی کارڈ اور عبارت پر نظر؛ درجہ C = صرف عنوانات اور مسائل۔ خالی چھوڑیں تو سب صفحے A شمار ہوں گے۔ سابقہ پرچوں کا نقشہ بننے کے بعد B اور C کے صفحات لکھیں۔",
    "3) 'منٹ فی صفحہ' ابتدائی اندازہ ہے۔ پہلے ہفتے کے بعد 'اصل منٹ فی صفحہ' دیکھ کر اسے بدل دیں۔",
    "4) 'پیچھے' لکھا آئے تو تین راستے ہیں: کچھ ابواب B/C میں ڈالیں، اس کتاب کے ہفتہ وار گھنٹے بڑھائیں، یا روزانہ گھنٹے بڑھائیں۔",
    "5) تیسیر مصطلح الحدیث پہلے (ہفتہ 1 تا 5)، شرح نخبۃ الفکر بعد میں (ہفتہ 6 تا 10) — نصاب کی ہدایت 1۔ ان دونوں کے 0.5 گھنٹے دس ہفتوں کا اوسط ہیں۔",
    "6) پرچہ اول کے نمبر نصاب کی ہدایت 3 سے: التبیان 30، تیسیر 25، شرح نخبہ 25، آئینہ قادیانیت 20۔ حفظ احادیث 15 نمبر کا ایک جز (ہدایت 2)۔",
]
for i, n in enumerate(notes):
    cell = q.cell(24 + i, 1, n)
    cell.font = f_note
    cell.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    q.merge_cells(start_row=24 + i, start_column=1, end_row=24 + i, end_column=23)
    q.row_dimensions[24 + i].height = 28

widths = {"A": 5, "B": 14, "C": 26, "D": 7, "E": 15, "F": 9, "G": 9, "H": 15, "I": 10, "J": 10,
          "K": 10, "L": 8, "M": 8, "N": 10, "O": 15, "P": 9, "Q": 9, "R": 9, "S": 20, "T": 11,
          "U": 9, "V": 16, "W": 16}
for k, v in widths.items():
    q.column_dimensions[k].width = v
q.freeze_panes = "D11"

q.conditional_formatting.add("S11:S20", FormulaRule(formula=['ISNUMBER(SEARCH("پیچھے",S11))'],
                             fill=PatternFill("solid", fgColor="F8D7DA"), font=Font(name=FONT, color="9C0006", bold=True)))
q.conditional_formatting.add("S11:S20", FormulaRule(formula=['OR(S11="ٹھیک",S11="مکمل ✓")'],
                             fill=PatternFill("solid", fgColor="D4EDDA"), font=Font(name=FONT, color="155724", bold=True)))
q.conditional_formatting.add("M11:M20", DataBarRule(start_type="num", start_value=0, end_type="num", end_value=1, color="63BE7B"))

# ---------------------------------------------------------------- log sheet
lg = wb.create_sheet("لاگ", 1)
lg.sheet_view.rightToLeft = True
lg["A1"] = "روزانہ لاگ"
lg["A1"].font = f_title
lg["A2"] = ("ہر نشست کے بعد ایک سطر: تاریخ، کتاب، کس صفحے سے کس صفحے تک، کتنے منٹ، کون سا طریقہ۔ "
            "'دہرائی' والی سطریں پہلے دور کی پیش رفت میں شمار نہیں ہوتیں بلکہ الگ گنی جاتی ہیں۔")
lg["A2"].font = f_sub
lg.merge_cells("A2:H2")
lg["A3"] = "مثال:  5 اکتوبر 2026 | ہدایہ جلد ثالث | 12 | 17 | (صفحات خود بنیں گے: 6) | 80 | تمہیدی کارڈ + عبارت | کتاب البیوع، خیار الشرط کی تمہید"
lg["A3"].font = f_note
lg.merge_cells("A3:H3")
head_row(lg, 5, ["تاریخ", "کتاب", "سے صفحہ", "تک صفحہ", "صفحات", "منٹ", "طریقہ", "نوٹ"])
for r in range(6, 1001):
    for col in (1, 2, 3, 4, 6, 7, 8):
        cell = lg.cell(r, col)
        cell.font = f_input
        cell.border = box
    lg.cell(r, 1).number_format = DATE_FMT
    e = lg.cell(r, 5, f'=IF(AND(ISNUMBER(C{r}),ISNUMBER(D{r})),MAX(0,D{r}-C{r}+1),"")')
    e.font = f_body
    e.border = box
    e.alignment = center
dv_book = DataValidation(type="list", formula1=f"={BOOK_LIST}", allow_blank=True)
dv_meth = DataValidation(type="list", formula1=f"={METHOD_LIST}", allow_blank=True)
dv_date = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True)
lg.add_data_validation(dv_book)
lg.add_data_validation(dv_meth)
lg.add_data_validation(dv_date)
dv_book.add("B6:B1000")
dv_meth.add("G6:G1000")
dv_date.add("A6:A1000")
for k, v in {"A": 13, "B": 26, "C": 9, "D": 9, "E": 8, "F": 8, "G": 20, "H": 40}.items():
    lg.column_dimensions[k].width = v
lg.freeze_panes = "A6"

# ---------------------------------------------------------------- weekly plan
wk = wb.create_sheet("ہفتہ وار نقشہ", 2)
wk.sheet_view.rightToLeft = True
wk["A1"] = "103 دن کا نقشہ — 28 ستمبر 2026 تا 8 جنوری 2027"
wk["A1"].font = f_title
wk["A2"] = ("ہدف فیصد = پہلے دور میں تمام کتابوں کے کل صفحات کا کتنا حصہ اس ہفتے کے آخر تک ہو جانا چاہیے۔ "
            "آخری دنوں کی ترتیب اس مفروضے پر ہے کہ پرچہ اول پہلے دن ہوگا؛ وفاق کا نظام الاوقات آنے پر ترتیب بدل لیں تاکہ پہلے پرچے کی دہرائی سب سے آخر میں ہو۔")
wk["A2"].font = f_sub
wk.merge_cells("A2:J2")
wk.row_dimensions[2].height = 30
head_row(wk, 4, ["ہفتہ / دن", "سے", "تک", "مرحلہ", "کیا کرنا ہے", "ہدف: پہلا دور (مجموعی)",
                 "اصل (مجموعی)", "ہدف گھنٹے", "اصل گھنٹے", "حالت"])

rows = []
d0 = dt.date(2026, 9, 28)
rows.append(("ہفتہ 0", d0, d0 + dt.timedelta(days=6), "تیاری",
             "صفحات درج کریں، ریکارڈنگز کی فہرست، سابقہ پرچوں کا نقشہ، ہدایہ کے ایک سبق پر تمہیدی کارڈ کا تجربہ، حفظ شروع", 0))
w1 = dt.date(2026, 10, 5)
focus = {
    1: "پہلا دور شروع — روزانہ ترتیب (نیچے)۔ تیسیر مصطلح الحدیث شروع",
    2: "پہلا دور — روزانہ ترتیب",
    3: "پہلا دور — ہفتے کے آخر میں پہلی بار 'حالت' دیکھ کر درجات A/B/C طے کریں",
    4: "پہلا دور — روزانہ ترتیب",
    5: "آدھا راستہ: 50% ہدف۔ تیسیر مکمل؛ منٹ فی صفحہ اور گھنٹے دوبارہ ترتیب دیں",
    6: "پہلا دور — شرح نخبۃ الفکر شروع (تیسیر کے بعد)",
    7: "پہلا دور — روزانہ ترتیب",
    8: "پہلا دور — آئینہ قادیانیت مکمل کرنے کا ہدف",
    9: "پہلا دور — روزانہ ترتیب؛ 38 احادیث مکمل یاد",
    10: "پہلا دور مکمل (100%)۔ جو رہ گیا اسے درجہ B کے طور پر نمٹائیں",
}
for n in range(1, 11):
    s = w1 + dt.timedelta(days=7 * (n - 1))
    rows.append((f"ہفتہ {n}", s, s + dt.timedelta(days=6), "پہلا دور", focus[n], n / 10))
rows.append(("ہفتہ 11", dt.date(2026, 12, 14), dt.date(2026, 12, 20), "دہرائی + پرچے",
             "ہدایہ ثالث 3 دن، ہدایہ رابع 3 دن (کتاب + اپنے حاشیے + تمہیدی نوٹ؛ ہر کتاب کے تیسرے دن ایک پورا پرچہ وقت لگا کر)، 1 دن کمزور ابواب", None))
rows.append(("ہفتہ 12", dt.date(2026, 12, 21), dt.date(2026, 12, 27), "دہرائی + پرچے",
             "مشکوٰۃ اول 3 دن، مشکوٰۃ دوم 3 دن (ہر ایک کا آخری دن پورا پرچہ)، 1 دن 38 احادیث مکمل", None))
rows.append(("ہفتہ 13", dt.date(2026, 12, 28), dt.date(2026, 12, 31), "دہرائی + پرچے",
             "بیضاوی 2 دن، پرچہ اول 2 دن (ہر ایک کا آخری دن پورا پرچہ)", None))
final_days = [
    "ہدایہ رابع: مشکل الفاظ + تمہیدی نوٹ + عبارت پر نظر",
    "ہدایہ ثالث: مشکل الفاظ + تمہیدی نوٹ + عبارت پر نظر",
    "مشکوٰۃ دوم",
    "مشکوٰۃ اول + 38 احادیث",
    "بیضاوی",
    "پرچہ اول: التبیان + آئینہ قادیانیت",
    "پرچہ اول: تیسیر + شرح نخبۃ الفکر",
    "ہلکا دن: پرچہ اول پر آخری نظر، 38 احادیث، جلدی سونا",
]
for i, t in enumerate(final_days):
    d = dt.date(2027, 1, 1) + dt.timedelta(days=i)
    rows.append((f"دن {i + 1}", d, d, "آخری دہرائی", t, None))
for i in range(6):
    d = dt.date(2027, 1, 9) + dt.timedelta(days=i)
    rows.append((f"امتحان {i + 1}", d, d, "امتحان",
                 "پرچہ — وفاق کے نظام الاوقات کے مطابق۔ شام: اگلے پرچے کے مشکل الفاظ اور تمہیدی نوٹ", None))

TOTAL_PAGES = "SUM('کوٹہ'!$G$11:$G$19)"
for i, (label, s, e, phase, what, target) in enumerate(rows):
    r = 5 + i
    calc(wk.cell(r, 1), label, bold=True)
    calc(wk.cell(r, 2), s, DATE_FMT)
    calc(wk.cell(r, 3), e, DATE_FMT)
    calc(wk.cell(r, 4), phase)
    c = wk.cell(r, 5, what)
    c.font = f_body
    c.border = box
    c.alignment = Alignment(horizontal="right", vertical="center", wrap_text=True)
    if target is not None:
        calc(wk.cell(r, 6), target, "0%")
        calc(wk.cell(r, 7),
             f'=IF(B{r}>TODAY(),"",IFERROR(SUMIFS({LOG_E},\'لاگ\'!$A$6:$A$1000,"<="&C{r},{LOG_G},"<>دہرائی")/{TOTAL_PAGES},""))',
             "0%")
        calc(wk.cell(r, 10),
             f'=IF(OR(B{r}>TODAY(),G{r}=""),"",IF(G{r}>=F{r},"ہدف پورا ✓",IF(C{r}>=TODAY(),"جاری","پیچھے")))')
    else:
        for col in (6, 7, 10):
            calc(wk.cell(r, col), "")
            wk.cell(r, col).fill = fill_grey
    if phase == "امتحان":
        calc(wk.cell(r, 8), "")
        calc(wk.cell(r, 9), "")
        for col in range(1, 11):
            wk.cell(r, col).fill = fill_light
    else:
        calc(wk.cell(r, 8), f"=(C{r}-B{r}+1)*'کوٹہ'!$F$4", "0.0")
        calc(wk.cell(r, 9), f"=SUMIFS({LOG_F},'لاگ'!$A$6:$A$1000,\">=\"&B{r},'لاگ'!$A$6:$A$1000,\"<=\"&C{r})/60", "0.0")
    wk.row_dimensions[r].height = 34
last_row = 5 + len(rows) - 1  # should be 32
assert last_row == 32, last_row
wk.conditional_formatting.add("J5:J18", FormulaRule(formula=['J5="پیچھے"'],
                              fill=PatternFill("solid", fgColor="F8D7DA"), font=Font(name=FONT, color="9C0006", bold=True)))
wk.conditional_formatting.add("J5:J18", FormulaRule(formula=['J5="ہدف پورا ✓"'],
                              fill=PatternFill("solid", fgColor="D4EDDA"), font=Font(name=FONT, color="155724", bold=True)))
wk.conditional_formatting.add("A5:J32", FormulaRule(formula=['AND($B5<=TODAY(),$C5>=TODAY())'],
                              fill=PatternFill("solid", fgColor="FFF4C2")))

wk["A35"] = "پہلے دور کی روزانہ ترتیب (3.5 گھنٹے: 90 + 75 + 30 منٹ + 15 منٹ حفظ)"
wk["A35"].font = f_label
wk.merge_cells("A35:F35")
head_row(wk, 36, ["WEEKDAY", "دن", "بلاک 1 (90 منٹ)", "بلاک 2 (75 منٹ)", "بلاک 3 (30 منٹ)", "روزانہ"])
rot = [
    (1, "اتوار", "ہدایہ جلد رابع", "مشکوٰۃ جلد دوم", "بیضاوی"),
    (2, "پیر", "ہدایہ جلد ثالث", "مشکوٰۃ جلد اول", "تیسیر (ہفتہ 1–5) / شرح نخبہ (ہفتہ 6–10)"),
    (3, "منگل", "ہدایہ جلد رابع", "مشکوٰۃ جلد دوم", "بیضاوی"),
    (4, "بدھ", "ہدایہ جلد ثالث", "مشکوٰۃ جلد اول", "تیسیر (ہفتہ 1–5) / شرح نخبہ (ہفتہ 6–10)"),
    (5, "جمعرات", "ہدایہ جلد رابع", "مشکوٰۃ جلد دوم", "بیضاوی"),
    (6, "جمعہ", "بیضاوی (60 منٹ) + التبیان (30 منٹ)", "آئینہ قادیانیت (30 منٹ) + ہفتہ وار جائزہ / سابقہ پرچہ (45 منٹ)", "لاگ اور 'حالت' دیکھ کر اگلے ہفتے کی ترتیب"),
    (7, "ہفتہ", "ہدایہ جلد ثالث", "مشکوٰۃ جلد اول", "التبیان"),
]
for i, row in enumerate(rot):
    r = 37 + i
    for col, v in enumerate(row, start=1):
        calc(wk.cell(r, col), v)
    calc(wk.cell(r, 6), "حفظ احادیث 15 منٹ" if row[0] != 6 else "حفظ 15 منٹ + " + row[4])
    wk.cell(r, 5).value = row[4] if row[0] != 6 else "—"
    wk.row_dimensions[r].height = 30
    wk.cell(r, 1).font = Font(name=FONT, size=8, color="999999")
wk["A45"] = ("جمعہ کو عموماً مدرسے کی چھٹی ہوتی ہے، اس لیے ہلکی کتابیں اور جائزہ اس دن رکھا ہے۔ آپ کی چھٹی کسی اور دن ہو تو صرف یہ جدول بدلیں۔ "
             "ہفتہ وار حساب: ہدایہ ثالث 4.5، رابع 4.5، مشکوٰۃ اول 3.75، دوم 3.75، بیضاوی 2.5، پرچہ اول 2.5، جائزہ 1.25، حفظ 1.75 = 24.5 گھنٹے۔")
wk["A45"].font = f_note
wk["A45"].alignment = Alignment(horizontal="right", wrap_text=True, vertical="top")
wk.merge_cells("A45:J46")
wk.row_dimensions[45].height = 30
for k, v in {"A": 10, "B": 12, "C": 12, "D": 14, "E": 60, "F": 14, "G": 12, "H": 10, "I": 10, "J": 12}.items():
    wk.column_dimensions[k].width = v
wk.freeze_panes = "A5"

# ---------------------------------------------------------------- hifz
hz = wb.create_sheet("حفظ احادیث", 3)
hz.sheet_view.rightToLeft = True
hz["A1"] = "حفظ احادیث — 38 احادیث از مشکوٰۃ المصابیح"
hz["A1"].font = f_title
hz["A2"] = ("نصاب ہدایت 2: مکمل سند ضروری نہیں، متن حدیث صحابی کے نام کے ساتھ یاد کریں۔ 15 نمبر کا ایک جز۔ "
            "یاد کرنے کی تاریخ لکھیں؛ دہرائی خود بتائے گی: 2 دن، پھر 7 دن، پھر 21 دن بعد۔")
hz["A2"].font = f_sub
hz.merge_cells("A2:K2")
hz["A3"] = ("کتاب الایمان کی احادیث کے ابتدائی الفاظ حوالے کے لیے لکھ دیے ہیں — اپنی کتاب سے ملا لیں۔ باقی خانے اپنی کتاب سے بھریں۔ "
            "باب حفظ اللسان کی 19 احادیث = کل 38 میں سے باقی (9 + 4 + 6 = 19)؛ اپنی کتاب میں گنتی دیکھ لیں۔")
hz["A3"].font = f_note
hz.merge_cells("A3:K3")
hz.row_dimensions[2].height = 28
hz.row_dimensions[3].height = 28
head_row(hz, 5, ["#", "حصہ", "حدیث", "صحابی", "ابتدائی الفاظ", "یاد کیا", "دہرائی 1", "دہرائی 2",
                 "دہرائی 3", "اگلی دہرائی", "حالت"])
iman = [
    (2, "عمر بن الخطابؓ", "بينما نحن عند رسول الله ﷺ ذات يوم إذ طلع علينا رجل…"),
    (3, "ابو ہریرہؓ", "(حدیث جبریل کی دوسری روایت)"),
    (4, "عبد اللہ بن عمرؓ", "بني الإسلام على خمس…"),
    (5, "ابو ہریرہؓ", "الإيمان بضع وسبعون شعبة…"),
    (6, "عبد اللہ بن عمروؓ", "المسلم من سلم المسلمون من لسانه ويده…"),
    (7, "انسؓ", "لا يؤمن أحدكم حتى أكون أحب إليه من والده وولده…"),
    (8, "انسؓ", "ثلاث من كن فيه وجد بهن حلاوة الإيمان…"),
    (9, "عباس بن عبد المطلبؓ", "ذاق طعم الإيمان من رضي بالله رباً…"),
    (10, "ابو ہریرہؓ", "والذي نفس محمد بيده لا يسمع بي أحد من هذه الأمة…"),
]
hrows = []
for n, s, t in iman:
    hrows.append(("کتاب الایمان", f"حدیث {n}", s, t))
for n in range(1, 5):
    hrows.append(("باب الاعتصام بالکتاب والسنۃ", f"ابتدائی {n}", "", ""))
for n in range(1, 7):
    hrows.append(("کتاب العلم", f"ابتدائی {n}", "", ""))
for n in range(1, 20):
    hrows.append(("باب حفظ اللسان والغیبۃ والشتم (الفصل الاول)", f"{n}", "", ""))
assert len(hrows) == 38
for i, (sec, num, sah, words) in enumerate(hrows):
    r = 6 + i
    calc(hz.cell(r, 1), i + 1)
    c = hz.cell(r, 2, sec)
    c.font = f_body
    c.border = box
    c.alignment = right
    calc(hz.cell(r, 3), num)
    inp(hz.cell(r, 4), sah if sah else None)
    inp(hz.cell(r, 5), words if words else None)
    hz.cell(r, 5).alignment = right
    for col in (6, 7, 8, 9):
        inp(hz.cell(r, col), None, DATE_FMT)
    calc(hz.cell(r, 10), f'=IF(F{r}="","",IF(G{r}="",F{r}+2,IF(H{r}="",G{r}+7,IF(I{r}="",H{r}+21,"پختہ ✓"))))', DATE_FMT)
    calc(hz.cell(r, 11), f'=IF(F{r}="","یاد کرنا ہے",IF(ISNUMBER(J{r}),IF(J{r}<=TODAY(),"آج دہرائیں","ٹھیک"),J{r}))')
hz.conditional_formatting.add("K6:K43", FormulaRule(formula=['K6="آج دہرائیں"'],
                              fill=PatternFill("solid", fgColor="FFE0B2"), font=Font(name=FONT, color="8A4B00", bold=True)))
hz.conditional_formatting.add("K6:K43", FormulaRule(formula=['K6="پختہ ✓"'],
                              fill=PatternFill("solid", fgColor="D4EDDA"), font=Font(name=FONT, color="155724", bold=True)))
for k, v in {"A": 5, "B": 30, "C": 10, "D": 18, "E": 46, "F": 12, "G": 12, "H": 12, "I": 12, "J": 13, "K": 13}.items():
    hz.column_dimensions[k].width = v
hz.freeze_panes = "A6"

# ---------------------------------------------------------------- past papers
pp = wb.create_sheet("سابقہ پرچے", 4)
pp.sheet_view.rightToLeft = True
pp["A1"] = "سابقہ پرچوں کا نقشہ (الجواب للموقوف علیہ — دس سالہ پرچے)"
pp["A1"].font = f_title
pp["A2"] = ("بائیں جدول میں ہر سوال کی ایک سطر۔ دائیں جدول میں باب کا نام لکھیں تو گنتی اور درجہ خود بنے گا: "
            "10 سال میں 3 یا زیادہ بار = A، 1–2 بار = B، کبھی نہیں = C۔ الجواب کی PDF بھیجیں تو یہ نقشہ Claude بھر دے گا۔")
pp["A2"].font = f_sub
pp.merge_cells("A2:L2")
pp.row_dimensions[2].height = 28
head_row(pp, 4, ["سال (ھ)", "پرچہ", "سوال / شق", "کتاب", "باب / مبحث", "صفحہ (بشریٰ)", "کیا پوچھا گیا"])
for c_, h in zip(range(9, 13), ["باب / مبحث", "کتاب", "کتنی بار", "درجہ"]):
    cell = pp.cell(4, c_, h)
    cell.font = f_head
    cell.fill = fill_head
    cell.alignment = center
    cell.border = box
for r in range(5, 505):
    for col in range(1, 8):
        cell = pp.cell(r, col)
        cell.font = f_input
        cell.border = box
dv_book2 = DataValidation(type="list", formula1=f"={BOOK_LIST}", allow_blank=True)
pp.add_data_validation(dv_book2)
dv_book2.add("D5:D504")
dv_book2.add("J5:J124")
for r in range(5, 125):
    inp(pp.cell(r, 9))
    pp.cell(r, 9).alignment = right
    inp(pp.cell(r, 10))
    calc(pp.cell(r, 11), f'=IF(I{r}="","",COUNTIFS($E$5:$E$504,I{r},$D$5:$D$504,J{r}))')
    calc(pp.cell(r, 12), f'=IF(I{r}="","",IF(K{r}>=3,"A",IF(K{r}>=1,"B","C")))')
pp.conditional_formatting.add("L5:L124", CellIsRule(operator="equal", formula=['"A"'],
                              fill=PatternFill("solid", fgColor="F8D7DA"), font=Font(name=FONT, bold=True, color="9C0006")))
pp.conditional_formatting.add("L5:L124", CellIsRule(operator="equal", formula=['"B"'],
                              fill=PatternFill("solid", fgColor="FFF4C2"), font=Font(name=FONT, bold=True)))
for k, v in {"A": 9, "B": 8, "C": 10, "D": 24, "E": 26, "F": 10, "G": 40, "H": 3, "I": 26, "J": 24, "K": 9, "L": 8}.items():
    pp.column_dimensions[k].width = v
pp.freeze_panes = "A5"

# ---------------------------------------------------------------- recordings
rc = wb.create_sheet("ریکارڈنگز", 5)
rc.sheet_view.rightToLeft = True
rc["A1"] = "ریکارڈنگز کی فہرست + صفحے سے ریکارڈنگ تلاش"
rc["A1"].font = f_title
rc["A2"] = ("ہر سبق کی ایک سطر: فائل کا نام اور وہ کن صفحات کا احاطہ کرتا ہے۔ اوپر کتاب اور صفحہ لکھیں تو متعلقہ سبق سامنے آ جائے گا۔ "
            "بعد میں AI سے بنا وقت کا اشاریہ (صفحہ ← منٹ:سیکنڈ) اسی فہرست پر بنے گا۔")
rc["A2"].font = f_sub
rc.merge_cells("A2:J2")
rc.row_dimensions[2].height = 28
rc["A4"] = "کتاب"
rc["A5"] = "صفحہ"
rc["A6"] = "ریکارڈنگ"
rc["A7"] = "استاد / سبق"
for c_ in ("A4", "A5", "A6", "A7"):
    rc[c_].font = f_label
    rc[c_].alignment = right
inp(rc["B4"])
inp(rc["B5"], None, "0")
dv_book3 = DataValidation(type="list", formula1=f"={BOOK_LIST}", allow_blank=True)
rc.add_data_validation(dv_book3)
LAST = 800
dv_book3.add("B4")
dv_book3.add(f"B11:B{LAST}")
ROWMATCH = (f'SUMPRODUCT(MAX(($B$11:$B${LAST}=$B$4)*($F$11:$F${LAST}<=$B$5)*($G$11:$G${LAST}>=$B$5)'
            f'*(ROW($B$11:$B${LAST})-10)))')
calc(rc["D6"], f'=IF(OR($B$4="",$B$5=""),0,{ROWMATCH})')
rc["D6"].font = Font(name=FONT, size=8, color="999999")
calc(rc["B6"], f'=IF(OR($B$4="",$B$5=""),"کتاب اور صفحہ لکھیں",IF($D$6=0,"اس صفحے کی ریکارڈنگ درج نہیں",INDEX($A$11:$A${LAST},$D$6)))', bold=True)
calc(rc["B7"], f'=IF($D$6=0,"",INDEX($D$11:$D${LAST},$D$6)&" — سبق "&INDEX($C$11:$C${LAST},$D$6))')
rc.merge_cells("B6:C6")
rc.merge_cells("B7:C7")
head_row(rc, 10, ["فائل کا نام", "کتاب", "سبق #", "استاد", "دورانیہ (منٹ)", "سے صفحہ", "تک صفحہ",
                  "ٹرانسکرپٹ", "تمہیدی کارڈ", "نوٹ"])
for r in range(11, LAST + 1):
    for col in range(1, 11):
        cell = rc.cell(r, col)
        cell.font = f_input
        cell.border = box
dv_yes = DataValidation(type="list", formula1='"ہاں,نہیں,جزوی"', allow_blank=True)
rc.add_data_validation(dv_yes)
dv_yes.add(f"H11:I{LAST}")

# --- recordings from the Colab inventory (recordings_inventory.csv), if present
import csv as _csv, re as _re
INV = os.environ.get("INVENTORY", "inventory_clean.csv")
inv_rows = []
if os.path.exists(INV) and not TEST:
    with open(INV, encoding="utf-8") as fh:
        rd = _csv.reader(fh)
        next(rd)
        for rw in rd:
            rw = [c.strip() for c in rw] + [""] * 12
            inv_rows.append(rw)
book_order = {b[1]: i for i, b in enumerate(books)}
inv_rows.sort(key=lambda rw: (book_order.get(rw[1], 99), rw[10] or "9999", rw[0]))
lesson_no = {}
for i, rw in enumerate(inv_rows):
    r = 11 + i
    fname, book, minutes, topic, date, path = rw[0], rw[1], rw[4], rw[9], rw[10], rw[11]
    if book:
        lesson_no[book] = lesson_no.get(book, 0) + 1
    m = _re.search(r"(كتاب|کتاب)\s*(.+?)(\s+\d+)?$", topic)
    note = ("كتاب " + m.group(2).strip()) if m else ""
    if "تكرار" in topic or "تکرار" in topic:
        note = (note + " · تکرار").strip(" ·")
    if not book:
        note = "اس سال کی کلاس — کتاب کی تصدیق کریں" if "input_audio" in path else "کتاب نامعلوم"
    rc.cell(r, 1, fname)
    rc.cell(r, 2, book or None)
    rc.cell(r, 3, lesson_no.get(book) if book else None)
    rc.cell(r, 5, float(minutes) if minutes else None)
    rc.cell(r, 8, rw[7] or "نہیں")
    rc.cell(r, 9, "نہیں")
    rc.cell(r, 10, note or None)
assert len(inv_rows) <= LAST - 10, len(inv_rows)
if inv_rows:
    rc["E4"] = f"{len(inv_rows)} ریکارڈنگز Colab کی فہرست سے (28 ستمبر 2026)۔ صفحات ٹرانسکرپٹ بننے کے بعد بھرے جائیں گے۔"
    rc["E4"].font = f_note
    rc.merge_cells("E4:J4")

# --- per-book summary (formulas over the table)
sum_head = ["کتاب", "فائلیں", "گھنٹے", "1.5x پر گھنٹے", "ٹرانسکرپٹ ہوئے"]
for j, h in enumerate(sum_head):
    c = rc.cell(1, 12 + j, h)
    c.font = f_head
    c.fill = fill_head
    c.alignment = center
    c.border = box
B_RNG, E_RNG, H_RNG = f"$B$11:$B${LAST}", f"$E$11:$E${LAST}", f"$H$11:$H${LAST}"
for i, b in enumerate(books):
    r = 2 + i
    c = rc.cell(r, 12, b[1])
    c.font = Font(name=FONT, size=10, bold=True)
    c.border = box
    c.alignment = right
    calc(rc.cell(r, 13), f'=COUNTIF({B_RNG},L{r})', "0")
    calc(rc.cell(r, 14), f'=SUMIF({B_RNG},L{r},{E_RNG})/60', "0.0")
    calc(rc.cell(r, 15), f'=N{r}/1.5', "0.0")
    calc(rc.cell(r, 16), f'=COUNTIFS({B_RNG},L{r},{H_RNG},"ہاں")', "0")
c = rc.cell(11, 12, "کل")
c.font = Font(name=FONT, size=10, bold=True)
c.border = box
c.alignment = right
calc(rc.cell(11, 13), f"=COUNTA($A$11:$A${LAST})", "0", bold=True)
calc(rc.cell(11, 14), f"=SUM({E_RNG})/60", "0.0", bold=True)
calc(rc.cell(11, 15), "=N11/1.5", "0.0", bold=True)
calc(rc.cell(11, 16), f'=COUNTIF({H_RNG},"ہاں")', "0", bold=True)
rc.cell(12, 12, "1.5x پر سننے کا وقت آپ کے کل ~360 گھنٹوں سے موازنہ کے لیے۔").font = f_note

for k, v in {"A": 40, "B": 24, "C": 8, "D": 16, "E": 10, "F": 9, "G": 9, "H": 11, "I": 11, "J": 30,
             "K": 3, "L": 26, "M": 9, "N": 9, "O": 12, "P": 13}.items():
    rc.column_dimensions[k].width = v
rc.freeze_panes = "A11"

# ---------------------------------------------------------------- test data
if TEST:
    today = dt.date.today()
    lg["A6"], lg["B6"], lg["C6"], lg["D6"], lg["F6"], lg["G6"] = today, "ہدایہ جلد ثالث", 12, 17, 80, "تمہیدی کارڈ + عبارت"
    lg["A7"], lg["B7"], lg["C7"], lg["D7"], lg["F7"], lg["G7"] = today, "ہدایہ جلد ثالث", 12, 14, 20, "دہرائی"
    q["E18"], q["F18"] = 11, 510
    q["H18"], q["I18"] = 200, 150
    rc["A11"], rc["B11"], rc["C11"], rc["D11"], rc["F11"], rc["G11"] = "H3_L01.mp3", "ہدایہ جلد ثالث", 1, "مفتی صاحب", 11, 16
    rc["A12"], rc["B12"], rc["C12"], rc["D12"], rc["F12"], rc["G12"] = "H3_L02.mp3", "ہدایہ جلد ثالث", 2, "مفتی صاحب", 17, 22
    rc["B4"], rc["B5"] = "ہدایہ جلد ثالث", 19
    hz["F6"] = today - dt.timedelta(days=3)
    pp["A5"], pp["D5"], pp["E5"] = "1446", "ہدایہ جلد ثالث", "خیار الشرط"
    pp["A6"], pp["D6"], pp["E6"] = "1445", "ہدایہ جلد ثالث", "خیار الشرط"
    pp["I5"], pp["J5"] = "خیار الشرط", "ہدایہ جلد ثالث"

wb.active = 0
wb.save(OUT)
print("saved", OUT)
